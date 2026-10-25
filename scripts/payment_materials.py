"""Immutable source snapshots, exact text excerpts and genuine PDF/image crops."""

import csv
import hashlib
import html
import io
import json
import shutil
from datetime import datetime, date
from pathlib import Path
from urllib.parse import quote

try:
    from .payment_evidence import digest
    from .import_pulse_statements import TX_RE, amount_to_cents
except ImportError:
    from payment_evidence import digest
    from import_pulse_statements import TX_RE, amount_to_cents


def checksum(path):
    digest = hashlib.sha256()
    with path.open('rb') as source:
        for block in iter(lambda: source.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def snapshot(source: Path, output: Path):
    key = checksum(source)
    relative = Path("materials/sources") / key / source.name
    dest = output / relative
    dest.parent.mkdir(parents=True, exist_ok=True)
    if not dest.exists():
        shutil.copyfile(source, dest)
    if checksum(dest) != key:
        raise ValueError("Source snapshot differs from its content hash")
    return {"path": relative.as_posix(), "sha256": key, "original_name": source.name}


def write_asset(output, relative, content):
    path = output / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    encoded = content.encode() if isinstance(content, str) else content
    if path.exists() and path.read_bytes() != encoded:
        raise ValueError(f"Refusing to overwrite immutable evidence {relative}")
    if not path.exists():
        path.write_bytes(encoded)
    return {"path": relative, "sha256": checksum(path)}


def excerpt_page(title, body, source, prefix, *, image=None, page=None):
    link = quote(prefix + source["path"], safe="/")
    if page:
        link += "#page=" + str(page)
    picture = f'<p><img src="{html.escape(image)}" style="max-width:100%;height:auto" alt="原始记录裁剪图"></p>' if image else ""
    return '<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>' + html.escape(title) + '</title><style>body{font:16px/1.7 -apple-system,system-ui,sans-serif;background:#f5f7f9;color:#243648;margin:0}main{max-width:1100px;margin:auto;padding:28px}table{border-collapse:collapse;width:100%;background:white}th,td{padding:10px;border:1px solid #dce3e9;text-align:left;vertical-align:top;white-space:pre-wrap;overflow-wrap:anywhere}pre{background:white;padding:18px;white-space:pre-wrap;overflow-wrap:anywhere}a{color:#245c80}</style><main><h1>' + html.escape(title) + '</h1><p>保留原始证据用于回忆与复核，解析结果不能替代原文。</p>' + picture + body + '<p><a href="' + html.escape(link) + '">打开保存的原始文件' + (f'（第{page}页）' if page else '') + '</a></p><p>来源：' + html.escape(source["original_name"]) + '</p></main></html>'


def csv_excerpts(path):
    data = path.read_bytes()
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("gb18030")
    physical = text.splitlines(keepends=True)
    reader = csv.reader(io.StringIO(text, newline=""))
    previous, result = 0, {}
    for ordinal, values in enumerate(reader, 1):
        result[ordinal] = {"values": values, "literal": "".join(physical[previous:reader.line_num]),
                           "line_from": previous + 1, "line_to": reader.line_num}
        previous = reader.line_num
    return result


def text_materials(root, output, documents):
    import openpyxl
    count = 0
    for document in documents:
        path = root / document.get("source_path", document["file_name"])
        source = snapshot(path, output)
        document["source_snapshot"] = source
        workbook = None
        if document["platform"] == "wechat":
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            header_index = next(i for i, row in enumerate(rows) if "交易时间" in row)
            header = list(rows[header_index])
        else:
            excerpts = csv_excerpts(path)
            anchor = "交易创建时间" if document["platform"] == "meituan" else "交易时间"
            header = next(r["values"] for r in excerpts.values() if anchor in r["values"])
        try:
            for row in document["observations"]:
                ordinal = row["source_row"]
                if workbook:
                    values = [v.isoformat(sep=" ") if isinstance(v, datetime) else v.isoformat() if isinstance(v, date) else v for v in rows[ordinal - 1]]
                    literal = "\t".join("" if v is None else str(v) for v in values)
                    locator = {"sheet": sheet.title, "row": ordinal, "columns": len(header), "kind": "xlsx_cells"}
                else:
                    original = excerpts[ordinal]
                    values, literal = original["values"], original["literal"]
                    locator = {"record_ordinal": ordinal, "physical_line_from": original["line_from"],
                               "physical_line_to": original["line_to"], "kind": "csv_record"}
                if len(values) != len(header):
                    raise ValueError(f"Evidence row column mismatch: {path.name}:{ordinal}")
                payload = {"source": source, "locator": locator,
                           "header": header, "values": values, "literal_text": literal,
                           "observation_id": row["observation_id"], "read_method": "native_export_cells"}
                key = digest("original-text-v1", source["sha256"], ordinal, payload)
                prefix = f"materials/records/{key}/"
                raw = write_asset(output, prefix + "record.json", json.dumps(payload, ensure_ascii=False, indent=2))
                text = write_asset(output, prefix + "record.txt", literal)
                table = '<table><thead><tr><th>原始字段</th><th>原始值</th></tr></thead><tbody>' + ''.join(
                    '<tr><th>' + html.escape(str(k or '')) + '</th><td>' + html.escape(str(v if v is not None else '')) + '</td></tr>'
                    for k, v in zip(header, values)) + '</tbody></table><h2>原始记录文本</h2><pre>' + html.escape(literal) + '</pre>'
                page_asset = write_asset(output, prefix + "record.html", excerpt_page(f"原始记录 · 第{ordinal}行", table, source, "../../../"))
                row["source_evidence"] = {"evidence_id": key, "kind": "text_excerpt", "source": source,
                                          "locator": locator, "preview": page_asset, "raw": raw, "text": text}
                count += 1
        finally:
            if workbook:
                workbook.close()
    return count


def bank_materials(root, output, bank):
    import pdfplumber
    import pypdfium2
    previous_path = output / "payment-evidence.json"
    previous = {r["source_evidence"]["evidence_id"]: r["source_evidence"]
                for r in json.loads(previous_path.read_text())["bank_records"]
                if r.get("source_evidence")} if previous_path.exists() else {}
    by_file = {}
    for row in bank:
        by_file.setdefault(row["statement_file"], []).append(row)
    count, unresolved = 0, []
    for name, records in by_file.items():
        source_path = root / "pulse_statements" / name
        if not source_path.exists():
            # The user may reorganize originals. The immutable, ledger-hashed
            # snapshot remains authoritative; never pick another same-name PDF.
            checksums = {r.get("statement_sha256") for r in records}
            if len(checksums) == 1 and None not in checksums:
                source_path = output / "materials/sources" / next(iter(checksums)) / name
        source = snapshot(source_path, output)
        if any(r.get("statement_sha256") and r["statement_sha256"] != source["sha256"] for r in records):
            raise ValueError("Statement source changed after bank import; reimport before attaching evidence")
        located = []
        with pdfplumber.open(source_path) as pdf:
            for page in pdf.pages:
                lines = page.extract_text_lines()
                if not any("Amount (CNY)" in line["text"] for line in lines):
                    continue
                for i, line in enumerate(lines):
                    match = TX_RE.match(line["text"])
                    if match:
                        data = match.groupdict()
                        located.append({"page": page.page_number, "page_size": [page.width, page.height],
                                        "data": data, "line": line, "following": lines[i + 1:i + 5], "used": False})
        document = pypdfium2.PdfDocument(source_path)
        render_cache = {}
        try:
            for row in sorted(records, key=lambda r: r["id"]):
                def agrees(position):
                    data = position["data"]
                    return (not position["used"] and data["post"] == row["post_date_raw"] and data["trans"] == row["trans_date_raw"]
                            and amount_to_cents(data["amount"]) == row["amount_cents"]
                            and bool(data["cr"]) == bool(row["is_credit"])
                            and " ".join(data["description"].split()) == " ".join(row["description"].split(" | ")[0].split()))
                position = next((p for p in located if agrees(p)), None)
                key = digest("bank-original-v2", source["sha256"], row["raw_line"], row["id"])
                prefix = f"materials/records/{key}/"
                image_asset, locator, missing_continuations = None, None, []
                if position:
                    position["used"] = True
                    line = position["line"]
                    bottom = line["bottom"]
                    for continuation in row["description"].split(" | ")[1:]:
                        next_line = next((l for l in position["following"] if " ".join(l["text"].split()) == " ".join(continuation.split())), None)
                        if next_line:
                            bottom = max(bottom, next_line["bottom"])
                        else:
                            missing_continuations.append(continuation)
                    w, h = position["page_size"]
                    box = [max(0, line["x0"] - 4), max(0, line["top"] - 0.15), min(w, line["x1"] + 4), min(h, bottom + 0.15)]
                    page_number = position["page"]
                    if page_number not in render_cache:
                        bitmap = document[page_number - 1].render(scale=2)
                        render_cache[page_number] = bitmap.to_pil().copy()
                        bitmap.close()
                    image = render_cache[page_number]
                    scale_x, scale_y = image.width / w, image.height / h
                    crop = image.crop((int(box[0] * scale_x), int(box[1] * scale_y), int(box[2] * scale_x + 1), int(box[3] * scale_y + 1)))
                    locator = {"kind": "pdf_region", "page": page_number, "page_size_points": [w, h],
                               "box_points_top_left": box, "crop_size_pixels": [crop.width, crop.height],
                               "matched_fields": ["post_date", "transaction_date", "description", "amount", "credit_flag"]}
                    locator["missing_continuations"] = missing_continuations
                    prior = previous.get(key)
                    if prior and prior.get("image"):
                        # Renderer upgrades can change PNG bytes. Reuse only the
                        # committed crop for this exact source and region.
                        if prior["source"] != source or prior["locator"] != locator:
                            raise ValueError("Retained bank crop source or region changed")
                        image_asset = prior["image"]
                        if image_asset["path"] != prefix + "record.png" or checksum(output / image_asset["path"]) != image_asset["sha256"]:
                            raise ValueError("Retained bank crop differs from its content hash")
                    else:
                        buffer = io.BytesIO(); crop.save(buffer, format="PNG")
                        image_asset = write_asset(output, prefix + "record.png", buffer.getvalue())
                else:
                    unresolved.append(row["id"])
                payload = {"source": source, "locator": locator,
                           "bank_record_id": row["id"], "literal_text": row["raw_line"], "read_method": "pdf_embedded_text"}
                raw = write_asset(output, prefix + "record.json", json.dumps(payload, ensure_ascii=False, indent=2))
                body = '<pre>' + html.escape(row["raw_line"]) + '</pre>'
                page_asset = write_asset(output, prefix + "record.html", excerpt_page("银行原始明细", body, source, "../../../",
                                  image="record.png" if image_asset else None, page=locator["page"] if locator else None))
                row["source_evidence"] = {"evidence_id": key, "kind": "pdf_crop" if image_asset else "pdf_text_unlocated",
                                          "source": source, "locator": locator, "preview": page_asset, "raw": raw,
                                          "image": image_asset, "crop_complete": bool(image_asset) and not missing_continuations}
                count += bool(image_asset)
        finally:
            document.close()
    return {"bank_crops": count, "unlocated_bank_records": unresolved}


def card_materials(root, output, jobs):
    results = {}
    for job in jobs:
        image_path = Path(job["imagePath"])
        if not image_path.exists():
            raise ValueError("Missing card evidence")
        source_video = job.get("sourceVideo") or job.get("source_video")
        if not source_video:
            raise ValueError("Missing original video provenance")
        source = snapshot(root / Path(source_video).name, output)
        image = snapshot(image_path, output)
        results[job["cardID"]] = {"kind": "order_card_crop", "source": source, "image": image,
                                    "crop_provenance": {k: v for k, v in job.items() if k not in {"expected", "imagePath"}},
                                    "actual_video_pts_verified": False}
    return results
